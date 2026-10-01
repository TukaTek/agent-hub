"""Builds Office files with the image's Python libraries and re-opens them, offline."""
import tempfile
import unittest
import zipfile
from pathlib import Path

import docx
import openpyxl
import pptx


class OfficeSmokeTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_xlsx_with_filter_merged_cells_and_link_reopens(self):
        path = self.dir / "pilot.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Region", "Units"])
        sheet.append(["North", 3])
        sheet.append(["South", 5])
        sheet.auto_filter.ref = "A1:B3"
        sheet.merge_cells("D1:E1")
        sheet["D1"] = "Summary"
        sheet["D2"].hyperlink = "https://example.com"
        workbook.save(path)

        reopened = openpyxl.load_workbook(path).active
        self.assertEqual(reopened.auto_filter.ref, "A1:B3")
        self.assertIn("D1:E1", [str(cells) for cells in reopened.merged_cells.ranges])
        self.assertEqual(reopened["B3"].value, 5)
        with zipfile.ZipFile(path) as archive:
            self.assertIsNone(archive.testzip())
            xml = archive.read("xl/worksheets/sheet1.xml").decode()
        # CT_Worksheet order: Excel rejects autoFilter after mergeCells or hyperlinks.
        self.assertLess(xml.index("<sheetData"), xml.index("<autoFilter"))
        self.assertLess(xml.index("<autoFilter"), xml.index("<mergeCells"))
        self.assertLess(xml.index("<mergeCells"), xml.index("<hyperlinks"))

    def test_docx_reopens(self):
        path = self.dir / "brief.docx"
        document = docx.Document()
        document.add_heading("Brief", 1)
        document.add_paragraph("Body")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Key"
        document.save(path)

        reopened = docx.Document(path)
        self.assertEqual([p.text for p in reopened.paragraphs], ["Brief", "Body"])
        self.assertEqual(reopened.tables[0].rows[0].cells[0].text, "Key")

    def test_pptx_reopens(self):
        path = self.dir / "deck.pptx"
        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Deck"
        slide.placeholders[1].text_frame.text = "Point"
        presentation.save(path)

        reopened = pptx.Presentation(path)
        self.assertEqual(len(reopened.slides), 1)
        self.assertEqual(reopened.slides[0].shapes.title.text, "Deck")


if __name__ == "__main__":
    unittest.main(verbosity=2)
